"""Inspect the SIH sample PPT structure."""
from pptx import Presentation
from pptx.util import Emu

src = r"C:\Users\Poloj\Downloads\Smartindia_Hackathon_132.pptx"
prs = Presentation(src)

import sys
sys.stdout.reconfigure(encoding='utf-8')

print(f"Slide width:  {prs.slide_width} EMU ({Emu(prs.slide_width).inches:.2f} in)")
print(f"Slide height: {prs.slide_height} EMU ({Emu(prs.slide_height).inches:.2f} in)")
print(f"Total slides: {len(prs.slides)}")
print(f"Slide masters: {len(prs.slide_masters)}")
print()

for sm_idx, sm in enumerate(prs.slide_masters):
    print(f"=== Slide master {sm_idx} ===")
    print(f"  Layouts: {len(sm.slide_layouts)}")
    for l_idx, layout in enumerate(sm.slide_layouts):
        print(f"    [{l_idx}] {layout.name}")
    print()

for i, slide in enumerate(prs.slides, 1):
    print(f"=== Slide {i} (layout: {slide.slide_layout.name}) ===")
    for sh in slide.shapes:
        if sh.has_text_frame:
            text = sh.text_frame.text.strip()
            if text:
                first = text.split('\n')[0][:80]
                print(f"  [{sh.shape_type}] '{sh.name}' pos=({Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f}) "
                      f"size=({Emu(sh.width).inches:.2f}x{Emu(sh.height).inches:.2f})")
                print(f"     first line: {first!r}")
        else:
            print(f"  [{sh.shape_type}] '{sh.name}' (no text frame) pos=({Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f})")
    print()
