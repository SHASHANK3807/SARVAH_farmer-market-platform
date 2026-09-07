"""
Inspect the existing Sarvah PPT to understand current content and layout.
"""
import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pptx import Presentation
from pptx.util import Emu

PPTX_PATH = r"C:\Users\Poloj\Downloads\sarvaha.pptx"

prs = Presentation(PPTX_PATH)

print(f"Total slides: {len(prs.slides)}")
print(f"Slide width: {prs.slide_width} EMU ({prs.slide_width / 914400:.1f} inches)")
print(f"Slide height: {prs.slide_height} EMU ({prs.slide_height / 914400:.1f} inches)")
print(f"Layouts: {len(prs.slide_layouts)}")
print()

for i, slide in enumerate(prs.slides, 1):
    print(f"=== SLIDE {i} ===")
    print(f"Layout: {slide.slide_layout.name}")
    print(f"Number of shapes: {len(slide.shapes)}")
    for j, shape in enumerate(slide.shapes):
        print(f"  Shape {j}: {shape.shape_type}, name='{shape.name}'")
        print(f"    Position: ({shape.left}, {shape.top}), Size: ({shape.width}, {shape.height})")
        if shape.has_text_frame:
            for k, para in enumerate(shape.text_frame.paragraphs):
                text = "".join(run.text for run in para.runs)
                if text.strip():
                    print(f"    Para {k}: '{text}'")
        if shape.has_table:
            print(f"    Table with {len(shape.table.rows)} rows x {len(shape.table.columns)} cols")
    print()
