"""Test: modify slide 1 only, save."""
from pptx import Presentation
import sys
sys.stdout.reconfigure(encoding='utf-8')

SRC = r"C:\Users\Poloj\Downloads\Smartindia_Hackathon_132.pptx"
DST = r"C:\Users\Poloj\Downloads\Mandi_Mitra_draft1.pptx"

prs = Presentation(SRC)


def set_paragraphs(shape, lines):
    tf = shape.text_frame
    paras = tf.paragraphs
    if not lines:
        for p in paras:
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
        return
    while len(paras) < len(lines):
        tf.add_paragraph()
    paras = tf.paragraphs
    for i, line in enumerate(lines):
        p = paras[i]
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        run = p.add_run()
        run.text = line


# Slide 1 modification
slide1 = prs.slides[0]
for sh in slide1.shapes:
    if sh.name == "Rectangle" and sh.has_text_frame:
        txt = sh.text_frame.text
        if "Problem Statement ID" in txt:
            print("Found target rectangle")
            new_lines = [
                "Idea Title -",
                "Mandi Mitra: Smart Price Discovery for Maharashtra Farmers",
                "Problem Statement ID -",
                "26132",
                "Theme -",
                "Agriculture, FoodTech & Rural Development",
                "PS Category - Software",
                "Team ID -",
                "Team Name (Registered on portal) -",
                "[Team Name TBD]",
            ]
            set_paragraphs(sh, new_lines)
            print("Modified")

print("Saving...")
prs.save(DST)
print("Saved.")
