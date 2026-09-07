"""Test: try modifying just one paragraph first."""
from pptx import Presentation
import sys
sys.stdout.reconfigure(encoding='utf-8')

SRC = r"C:\Users\Poloj\Downloads\Smartindia_Hackathon_132.pptx"
DST = r"C:\Users\Poloj\Downloads\Mandi_Mitra_draft1.pptx"

prs = Presentation(SRC)
slide1 = prs.slides[0]

# Find the target rectangle
for sh in slide1.shapes:
    if sh.name == "Rectangle" and sh.has_text_frame:
        txt = sh.text_frame.text
        if "Problem Statement ID" in txt:
            print("Found. Inspecting paragraphs:")
            for i, p in enumerate(sh.text_frame.paragraphs):
                print(f"  p{i}: {[r.text for r in p.runs]}")
            print(f"Total paragraphs: {len(sh.text_frame.paragraphs)}")
            break

# Now test: just modify first run, save
print("\nModifying first run only...")
for sh in slide1.shapes:
    if sh.name == "Rectangle" and sh.has_text_frame:
        txt = sh.text_frame.text
        if "Problem Statement ID" in txt:
            sh.text_frame.paragraphs[0].runs[0].text = "TEST"
            break

print("Saving...")
prs.save(DST)
print("Saved.")
