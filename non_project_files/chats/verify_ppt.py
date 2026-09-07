"""
Verify the updated PPT content
"""
import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pptx import Presentation

PPTX_PATH = r"C:\Users\Poloj\Downloads\sarvaha_new.pptx"

prs = Presentation(PPTX_PATH)

print(f"Total slides: {len(prs.slides)}")
print()

for i, slide in enumerate(prs.slides, 1):
    print(f"=== SLIDE {i} ===")
    for j, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            for k, para in enumerate(shape.text_frame.paragraphs):
                text = "".join(run.text for run in para.runs)
                if text.strip():
                    print(f"  [S{j+1}.P{k}]: {text[:120]}{'...' if len(text)>120 else ''}")
    print()
