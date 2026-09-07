"""Dump full text of each slide."""
from pptx import Presentation
import sys
sys.stdout.reconfigure(encoding='utf-8')

src = r"C:\Users\Poloj\Downloads\Smartindia_Hackathon_132.pptx"
prs = Presentation(src)

for i, slide in enumerate(prs.slides, 1):
    print(f"\n========== SLIDE {i} ==========")
    for sh in slide.shapes:
        if sh.has_text_frame:
            print(f"\n--- Shape: {sh.name} ---")
            for p_idx, para in enumerate(sh.text_frame.paragraphs):
                for r in para.runs:
                    print(f"  [p{p_idx}] {r.text!r}")
